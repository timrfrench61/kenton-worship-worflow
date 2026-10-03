"""Real remote-script entry point with fake service/HTTP controls, no network."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]


class DeploymentTests(unittest.TestCase):
    def test_workbook_log_is_idempotent_and_preserves_planning(self):
        sys.path.insert(0, str(ROOT/'scripts'))
        from _website_release import log_workbook
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root/'recent-logs.xlsx'
            book = openpyxl.Workbook()
            book.active.title = 'planner'
            book.active['A1'] = '=1+2'
            book.save(path)
            book.close()
            folder = root/'release'
            folder.mkdir()
            settings = {'public_url': 'https://example.test', 'maintenance_workbook': {
                'enabled': True, 'path': str(path), 'sheet': 'Website Maintenance'}}
            receipt = {'release_id': 'r1', 'status': 'deployed', 'note': '=literal description'}
            self.assertEqual(log_workbook(root, settings, folder, receipt), 'logged')
            self.assertEqual(log_workbook(root, settings, folder, receipt), 'already_logged')
            book = openpyxl.load_workbook(path)
            self.assertEqual(book['planner']['A1'].value, '=1+2')
            self.assertEqual(book['Website Maintenance'].max_row, 2)
            self.assertEqual(book['Website Maintenance']['E2'].data_type, 's')
            book.close()
            self.assertTrue((folder/'recent-logs-before.xlsx').exists())

    def exercise(self, fail=False, corrupt=False):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            live = root / 'server/site/live'
            live.mkdir(parents=True)
            (live/'old.txt').write_text('previous')
            (live/'appsettings.json').write_text('server configuration')
            release = root/'releases/example'
            release.mkdir(parents=True)
            payload = root/'payload'
            payload.mkdir()
            (payload/'new.txt').write_text('new content')
            (payload/'appsettings.json').write_text('development configuration')
            files = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in payload.iterdir()}
            if corrupt:
                files['new.txt'] = 'incorrect'
            with tarfile.open(release/'payload.tar.gz', 'w:gz') as archive:
                for p in payload.iterdir():
                    archive.add(p, arcname=p.name)
            (release/'release.json').write_text(json.dumps(dict(release_id='example', remote_directory=str(live),
                service_name='kenton-website', preserve_files=['appsettings.json'], files=files, health_url='https://example.test')))
            shutil.copyfile(ROOT/'scripts/_website_remote.py',root/'activate.py')
            (root/'sitecustomize.py').write_text('''
import io, os, subprocess, time, urllib.request
from pathlib import Path
starts = 0
def run(command, **kwargs):
    global starts
    with open(os.environ['CALLS'], 'a') as f: f.write(' '.join(command)+'\\n')
    if command[-2] == 'start' and '-l' not in command: starts += 1
    return subprocess.CompletedProcess(command, 0)
def open_url(*args, **kwargs):
    if os.environ.get('FAIL') and starts < 2: raise OSError('simulated unhealthy new release')
    result = io.BytesIO(b'ok'); result.status = 200; return result
subprocess.run = run
urllib.request.urlopen = open_url
time.sleep = lambda _: None
''')
            env = dict(os.environ, PYTHONPATH=str(root), CALLS=str(root/'calls.txt'))
            if fail:
                env['FAIL'] = '1'
            result = subprocess.run([sys.executable,str(root/'activate.py'),str(release)],env=env,capture_output=True,text=True)
            if corrupt:
                self.assertNotEqual(result.returncode, 0)
                self.assertEqual((live/'old.txt').read_text(), 'previous')
                self.assertNotIn('sudo -n /usr/bin/systemctl stop', (root/'calls.txt').read_text())
            elif fail:
                self.assertNotEqual(result.returncode, 0)
                self.assertEqual(json.loads((release/'remote-receipt.json').read_text())['status'], 'rolled_back')
                self.assertTrue((live/'old.txt').exists())
                self.assertFalse((live/'new.txt').exists())
            else:
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertTrue((live/'new.txt').exists())
                self.assertTrue((release/'previous/old.txt').exists())
            self.assertEqual((live/'appsettings.json').read_text(), 'server configuration')

    def test_activation_preserves_settings_and_backup(self):
        self.exercise()

    def test_failed_health_restores_previous_release(self):
        self.exercise(fail=True)

    def test_hash_failure_does_not_stop_live_service(self):
        self.exercise(corrupt=True)
