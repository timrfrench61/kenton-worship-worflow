import json
from datetime import datetime
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl
from kenton_workflow.planner import LABELS
from kenton_workflow.automation import plain_plan


class PlannerPrayerTests(unittest.TestCase):
    def test_cli_inserted_prayers_and_legacy(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            planner, catalog = root / 'recent-logs.xlsx', root / 'Song-Lists.xlsx'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = 'planner'
            sheet.append(['Date', datetime(2026, 10, 4)])
            for row, label in LABELS.items():
                sheet.cell(row, 1, label)
            sheet['B2'] = 'Matthew 5:8'
            sheet['B15'] = 'MOV'
            book.save(planner)
            songs = openpyxl.Workbook()
            for name in ('Praise', 'Hymns'):
                songs.create_sheet(name).append(['Name'])
            songs.save(catalog)
            songs.close()

            def read():
                result = subprocess.run([sys.executable, '-m', 'kenton_workflow', 'plan',
                                         '--date', '2026-10-04', '--planner', str(planner),
                                         '--catalog', str(catalog)], capture_output=True, text=True)
                self.assertIn(result.returncode, (0, 2), result.stderr)
                return json.loads(result.stdout)

            self.assertFalse(plain_plan(read())['morning']['has_prayer_text_row'])
            sheet.insert_rows(20)
            sheet['A20'], sheet['B20'] = 'PofC-text', 'Evening supplied prayer.'
            sheet.insert_rows(7)
            sheet['A7'], sheet['B7'] = 'PofC-text', 'Morning supplied prayer.'
            book.save(planner)
            raw = read()
            plan = plain_plan(raw)
            self.assertEqual(plan['morning']['prayer_of_confession_text'], 'Morning supplied prayer.')
            self.assertEqual(plan['evening']['prayer_of_confession_text'], 'Evening supplied prayer.')
            self.assertEqual(raw['evening']['prayer_of_confession_text']['source']['cell'], 'B21')
            self.assertEqual(raw['unmapped'], [])
            self.assertEqual(plan['evening']['sermon'], 'MOV')
            sheet['B7'] = '=1+1'
            book.save(planner)
            raw = read()
            self.assertIsNone(plain_plan(raw)['morning']['prayer_of_confession_text'])
            self.assertTrue(any('needs explicit prayer text' in issue for issue in raw['issues']))
            self.assertEqual(plain_plan(raw)['evening']['prayer_of_confession_text'], 'Evening supplied prayer.')
            book.close()
