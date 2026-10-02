import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from kenton_workflow.study_research import load_ranked
from kenton_workflow.handouts import generate
from test_handouts import template
from docx import Document
from docx.shared import Pt


class StudyResearchTests(unittest.TestCase):
    def test_source_validation_selection_and_template_mapping(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            cache = root / 'work/research/cache'
            cache.mkdir(parents=True)
            bible = 'test-niv'
            (root / 'application.json').write_text(json.dumps({'bible': {'bible_id': bible},
                'word_study': {'cache_directory': 'work/research/cache', 'body_font_size': 13}}))
            service = {'kind': 'word-study', 'sermon': 'Matt 5:8', 'topic': 'Test study', 'study_words': ['pure']}
            report = {'status': 'gemini_ranked_draft_not_approved', 'complete': True,
                      'passage': 'Matthew 5:8', 'ranking': {'created_utc': '2026-10-01T12:00:00+00:00'},
                      'bible': {'data': {'id': bible, 'abbreviation': 'NIV11', 'language': {'id': 'eng'}}},
                      'main_passage': {'data': {'passages': [{'id': 'MAT.5.8'}]}}, 'words': [{'word': 'pure'}]}
            for testament, book in [('old_testament', 'GEN'), ('new_testament', 'MAT')]:
                verses = []
                for number in range(1, 7):
                    identity = f'{book}.1.{number}'
                    ref = f'{book} 1:{number}'
                    text = f'Synthetic full verse {number}.'
                    url = f'https://rest.api.bible/v1/bibles/{bible}/verses/{identity}'
                    source = {'url': url, 'data': {'id': identity, 'bookId': book, 'bibleId': bible, 'reference': ref, 'content': text}}
                    (cache / (hashlib.sha256(url.encode()).hexdigest() + '.json')).write_text(json.dumps(source))
                    verses.append({'id': identity, 'reference': ref, 'text': text, 'source': source})
                report['words'][0][testament] = {'verses': verses}
            path = root / 'work/research/2026-10-04-ranked.json'
            path.write_text(json.dumps(report))
            selected, content, texts = load_ranked(root, '2026-10-04', service)
            self.assertEqual(selected, path)
            self.assertEqual(content['words'][0]['word'], 'Pure')
            self.assertEqual(len(content['words'][0]['old_testament']), 3)
            self.assertEqual(content['words'][0]['old_testament_see_also'], ['GEN 1:4', 'GEN 1:5', 'GEN 1:6'])
            self.assertIsNone(load_ranked(root, '2026-10-11', service))
            layout = root / 'template.docx'
            template(layout, 'word-study')
            document = Document(layout)
            for cell in document.tables[0].add_row().cells:
                cell.text = 'See also [References]'
            for row in document.tables[0].rows[1:]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(12)
            document.save(layout)
            content['passage_text'] = 'Synthetic main passage.'
            for word in content['words']:
                for testament in ('old_testament', 'new_testament'):
                    for item in word[testament]:
                        from kenton_workflow.handouts import reference
                        item['text'] = texts[reference(item['reference'])]
            output = root / 'output.docx'
            generate(service, {'_handout_source': layout, '_verified_word_content': content,
                               '_preserve_template_format': True}, output)
            doc = Document(output)
            self.assertEqual('See also GEN 1:4; GEN 1:5; GEN 1:6 (NIV)', doc.tables[0].cell(2, 0).text)
            self.assertNotIn('See also', doc.tables[0].cell(1, 0).text)
            self.assertNotIn('[References]', doc.tables[0].cell(2, 0).text)
            self.assertTrue(all(r.font.size.pt == 12 for p in doc.tables[0].cell(1, 0).paragraphs
                                for r in p.runs if r.text))
            revised = Document(layout)
            title = revised.paragraphs[0]._element
            title.getparent().remove(title)
            old_table = revised.tables[0]._element
            stacked = revised.add_table(rows=5, cols=1)
            for cell, value in zip([r.cells[0] for r in stacked.rows],
                                   ['OLD TESTAMENT', 'Reference — Text\nSee also [Reference]',
                                    'NEW TESTAMENT', 'Reference — Text', 'See also [Reference]']):
                cell.text = value
            old_table.addprevious(stacked._element)
            old_table.getparent().remove(old_table)
            revised.save(layout)
            generate(service, {'_handout_source': layout, '_verified_word_content': content,
                               '_preserve_template_format': True}, output)
            revised_output = Document(output)
            self.assertEqual(revised_output.paragraphs[0].text, 'A Word Study on Matt 5:8')
            self.assertEqual(revised_output.paragraphs[1].text, 'Synthetic main passage.')
            self.assertNotIn('Test study', [p.text for p in revised_output.paragraphs])
            self.assertEqual(len(revised_output.tables[0].columns), 1)
            config = json.loads((root / 'application.json').read_text())
            config['word_study']['handout_full_verses_per_testament_per_word'] = 1
            (root / 'application.json').write_text(json.dumps(config))
            _, reduced, _ = load_ranked(root, '2026-10-04', service)
            self.assertEqual(len(reduced['words'][0]['old_testament']), 1)
            self.assertEqual(len(reduced['words'][0]['old_testament_see_also']), 5)
            from copy import deepcopy
            config['word_study']['handout_full_verses_by_word_count'] = {'3': 3, '4': 2}
            (root / 'application.json').write_text(json.dumps(config))
            for number, expected in ((3, 3), (4, 2)):
                expanded = deepcopy(report)
                expanded['words'] = [dict(deepcopy(report['words'][0]), word=f'Word{i}') for i in range(number)]
                path.write_text(json.dumps(expanded))
                matching = dict(service, study_words=[w['word'] for w in expanded['words']])
                _, selected_content, _ = load_ranked(root, '2026-10-04', matching)
                for word in selected_content['words']:
                    for testament in ('old_testament', 'new_testament'):
                        self.assertEqual(len(word[testament]), expected)
                        self.assertEqual(len(word[testament + '_see_also']), 6 - expected)
            report['words'][0]['old_testament']['verses'][0]['text'] = 'Altered text'
            path.write_text(json.dumps(report))
            with self.assertRaisesRegex(ValueError, 'differs'):
                load_ranked(root, '2026-10-04', service)
