from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import RGBColor
from kenton_workflow.automation import bulletin, archive_older_generated


class BulletinTemplateTests(unittest.TestCase):
    def test_scripture_survives_blank_week_and_missing_slot_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / 'template.docx'
            doc = Document()
            doc.add_paragraph('September 20, 2026')
            doc.add_paragraph('CALL TO WORSHIP')
            table = doc.add_table(rows=2, cols=2)
            for row, speaker in zip(table.rows, ('Leader', 'People')):
                row.cells[0].text = speaker
                row.cells[1].text = 'Reading'
            for text in ('PRAISE MUSIC', 'Song', 'INVOCATION', 'HYMN - Hymn',
                         'SCRIPTURE - Example 1:1', 'PRAYER OF CONFESSION', 'Prayer',
                         'Reference', 'SILENT PRAYER', 'SERMON - Example', 'Title', 'COMMUNION'):
                doc.add_paragraph(text)
            heading = next(p for p in doc.paragraphs if p.text == 'PRAYER OF CONFESSION')
            heading.clear()
            heading.add_run('\ufffd').font.color.rgb = RGBColor.from_string('000000')
            heading.add_run('PRAYER OF CONFESSION').font.color.rgb = RGBColor.from_string('2F5496')
            doc.save(template)
            service = dict(praise_songs=['Song'], hymns=['Hymn'], call_to_worship='Example 1:1',
                           prayer_of_confession='Example 2:1', additional_reading='',
                           kind='word-study', sermon='Example 3:1')
            content = dict(prayer='Supplied prayer.', call_lines=['Line one', 'Line two'], display_title='New title')
            blank = root / 'blank.docx'
            bulletin(template, service, content, '2026-09-27', blank)
            service['additional_reading'] = 'Exodus 33:12-23'
            filled = root / 'filled.docx'
            bulletin(blank, service, content, '2026-10-04', filled)
            texts = [p.text for p in Document(filled).paragraphs]
            self.assertIn('SCRIPTURE — Exodus 33:12-23', texts)
            self.assertIn('COMMUNION', texts)
            heading = next(p for p in Document(filled).paragraphs if p.text.startswith('PRAYER OF CONFESSION'))
            self.assertEqual(heading.text, 'PRAYER OF CONFESSION (UNISON)')
            self.assertEqual(str(heading.runs[0].font.color.rgb), '2F5496')
            doc = Document(template)
            for p in doc.paragraphs:
                if p.text.startswith('SCRIPTURE'):
                    p._element.getparent().remove(p._element)
            doc.save(template)
            with self.assertRaisesRegex(ValueError, 'Scripture reference cannot be placed'):
                bulletin(template, service, content, '2026-10-04', filled)

    def test_archive_keeps_current_and_unrecognized_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name in ('desktop', 'output'):
                folder = root / 'work' / name
                folder.mkdir(parents=True)
                for file in ('2026-09-27-morning-bulletin.docx', '2026-10-04-morning-bulletin.docx', 'my-notes.docx'):
                    (folder / file).write_bytes(b'preserve')
            archive_older_generated(root, '2026-10-04')
            self.assertEqual(len(list((root / 'work/_archive').rglob('2026-09-27-morning-bulletin.docx'))), 2)
            self.assertTrue((root / 'work/output/2026-10-04-morning-bulletin.docx').exists())
            self.assertTrue((root / 'work/desktop/my-notes.docx').exists())
