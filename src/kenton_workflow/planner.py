"""Read-only planner extraction with explicit layout validation and provenance."""
from datetime import date, datetime
import hashlib
import json
import os
import re
from pathlib import Path
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter

from kenton_workflow.sync import checked_path

LABELS = {
    2: 'MSermon', 3: 'MTopic', 4: 'Otread', 5: 'C2W', 6: 'PofC',
    **{7+i: f'MP{i+1}' for i in range(4)},
    **{11+i: f'MH{i+1}' for i in range(3)},
    14: 'Otread', 15: 'ESermon', 16: 'ETopic', 18: 'C2W', 19: 'PofC',
    **{20+i: f'EP{i+1}' for i in range(4)},
    **{24+i: f'EH{i+1}' for i in range(9)},
}


def praise_title(value):
    """Ignore a final page reference for matching; retain arrangement subtitles."""
    return re.sub(r'\s*\([0-9IVXLCDMil]+(?:\s*[-–,]\s*[0-9IVXLCDMil]+)*\)\s*$', '', value).strip()


def song_key(value, tab):
    title = praise_title(value) if tab == 'Praise' else value.strip()
    return ' '.join(title.casefold().split())


def fingerprint(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def extract_week(planner_path, catalog_path, requested_date):
    day = date.fromisoformat(requested_date)
    if day.weekday() != 6:
        raise ValueError('Select a Sunday in YYYY-MM-DD format.')
    paths = [Path(planner_path), Path(catalog_path)]
    hashes = [fingerprint(p) for p in paths]
    book = openpyxl.load_workbook(paths[0], read_only=True, data_only=False)
    catalog = None
    try:
        catalog = openpyxl.load_workbook(paths[1], read_only=True, data_only=False)
        if 'planner' not in book.sheetnames:
            raise ValueError('Missing planner sheet.')
        sheet = book['planner']
        if sheet['A1'].value != 'Date':
            raise ValueError('Expected Date at planner!A1.')
        columns = [cell.column for cell in next(sheet.iter_rows())
                   if isinstance(cell.value, (date, datetime))
                   and cell.value.strftime('%Y-%m-%d') == day.isoformat()]
        if len(columns) != 1:
            raise ValueError(f'Expected one date column for {day}; found {len(columns)}.')
        column = get_column_letter(columns[0])
        # Resolve named fields after inserted rows rather than freezing row numbers.
        label_rows = {}
        for row in range(2, sheet.max_row + 1):
            label = sheet.cell(row, 1).value
            if isinstance(label, str):
                label_rows.setdefault(label.strip().casefold(), []).append(row)
        positions, occurrences = {}, {}
        for logical_row, label in LABELS.items():
            key = label.casefold()
            occurrence = occurrences.get(key, 0)
            occurrences[key] = occurrence + 1
            matches = label_rows.get(key, [])
            if occurrence >= len(matches):
                if key == 'otread':
                    positions[logical_row] = None
                    continue
                raise ValueError(f'Missing planner field {label!r}.')
            positions[logical_row] = matches[occurrence]
        issues = []
        def value(row):
            row = positions.get(row, row)
            if row is None:
                return {'value': None, 'source': {'workbook': paths[0].name, 'sheet': 'planner', 'cell': None}}
            cell = sheet[f'{column}{row}']
            raw = cell.value
            if cell.data_type == 'f':
                issues.append(f'{column}{row}: formula requires explicit evaluation; not imported as a result')
            if raw is not None and not isinstance(raw, str):
                issues.append(f'{column}{row}: expected text or blank')
            return {'value': raw if raw is None or isinstance(raw, str) else str(raw),
                    'source': {'workbook': paths[0].name, 'sheet': 'planner', 'cell': f'{column}{row}'}}
        indexes = {}
        for tab in ('Praise', 'Hymns'):
            if tab not in catalog.sheetnames or catalog[tab]['A1'].value != 'Name':
                raise ValueError(f'Expected {tab}!A1 header Name in song catalog.')
            index = {}
            for row in catalog[tab].iter_rows(min_row=2, max_col=1):
                cell = row[0]
                if isinstance(cell.value, str) and cell.data_type != 'f':
                    index.setdefault(song_key(cell.value, tab), []).append(cell.coordinate)
            indexes[tab] = index
        def songs(rows, tab):
            result = []
            for slot, row in enumerate(rows, 1):
                item = {'slot': slot, **value(row)}
                matches = indexes[tab].get(song_key(item['value'], tab), []) if item['value'] else []
                item['catalog_matches'] = [{'workbook': paths[1].name, 'sheet': tab, 'cell': c} for c in matches]
                if item['value'] and len(matches) != 1:
                    issues.append(f'{column}{row}: expected one {tab} name match; found {len(matches)}')
                result.append(item)
            return result
        def service(sermon, topic, call, confession, praise, hymns):
            source_text = value(sermon)
            return {'scheduled': None, 'sermon': {'source_text': source_text, 'topic': value(topic)},
                    'program_kind': 'video' if source_text['value'] == 'MOV' else None,
                    'call_to_worship': value(call), 'prayer_of_confession': value(confession),
                    'praise_songs': songs(praise, 'Praise'), 'hymns': songs(hymns, 'Hymns')}
        morning = service(2, 3, 5, 6, range(7, 11), range(11, 14))
        evening = service(15, 16, 18, 19, range(20, 24), range(24, 33))
        additional = [{'label': 'Otread', **value(row)} for row in (4, 14)]
        for entry in additional:
            if entry['value']:
                issues.append(f"{entry['source']['cell']}: additional reading service assignment needs review")
        unmapped = []
        for service_record in (morning, evening):
            service_record['study_words'] = {'value': None, 'source': None}
        for row in range(2, sheet.max_row + 1):
            if row not in positions.values() and sheet[f'{column}{row}'].value is not None:
                raw_value = sheet[f'{column}{row}'].value
                label_key = re.sub(r'[^a-z]', '', str(sheet[f'A{row}'].value).casefold())
                if label_key in ('studywords', 'mstudywords', 'estudywords'):
                    target = evening if label_key == 'estudywords' or row >= positions[15] else morning
                    target['study_words'] = {'value': raw_value,
                        'source': {'workbook': paths[0].name, 'sheet': 'planner', 'cell': f'{column}{row}'}}
                    continue
                unmapped.append({'label': sheet[f'A{row}'].value, 'value': raw_value,
                                 'source': {'workbook': paths[0].name, 'sheet': 'planner', 'cell': f'{column}{row}'}})
                issues.append(f'{column}{row}: populated unmapped row')
        if not any(sheet[f'{column}{r}'].value is not None for r in positions.values() if r is not None):
            issues.append('Selected column contains a date only; no service content.')
        result = {'schema_version': 1, 'date': day.isoformat(), 'status': 'draft', 'approved': False,
                  'sources': [{'file': str(p), 'sha256': h} for p, h in zip(paths, hashes)],
                  'date_cell': f'planner!{column}1', 'morning': morning, 'evening': evening,
                  'unassigned_readings': additional, 'unmapped': unmapped, 'issues': issues,
                  'limitations': ['Service scheduled flags are unset.', 'Catalog matches do not resolve chord or media files.',
                                  'Reading references are preserved; full liturgical text is not imported.',
                                  'Draft extraction does not establish approval or production completeness.']}
    finally:
        book.close()
        if catalog is not None:
            catalog.close()
    if hashes != [fingerprint(p) for p in paths]:
        raise ValueError('Source workbook changed during import; save and retry.')
    return result


def write_snapshot(result, destination):
    """Publish a new snapshot atomically; never replace a previous snapshot."""
    if result['issues']:
        raise ValueError('Resolve preview issues before writing a snapshot.')
    destination = checked_path(Path(destination))
    if destination.exists():
        raise ValueError(f'Snapshot already exists: {destination}. Preserve it and choose a new filename.')
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=destination.parent, delete=False) as stream:
            temporary = Path(stream.name)
            json.dump(result, stream, ensure_ascii=False, indent=2)
            stream.write('\n')
        os.link(temporary, destination)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
