"""Local release preparation, deployment receipts, and optional workbook logging."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shlex
import shutil
import subprocess
import tarfile
import tempfile
from urllib.parse import urlparse


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def history(root, config, item):
    path = (root / config['desktop_history']).resolve()
    if (root / 'work').resolve() not in path.parents:
        raise ValueError('Desktop history must be under work/.')
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('a', encoding='utf-8') as stream:
        stream.write(f"{datetime.now(timezone.utc).isoformat()}  {item['release_id']}  {item['status']}  {item.get('note', '')}\n")


def prepare(root, config, note):
    settings = config['deployment']
    output = (root / settings['history_directory']).resolve()
    if (root / 'work').resolve() not in output.parents:
        raise ValueError('Release history must be under work/.')
    release_id = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    folder = output / release_id
    folder.mkdir(parents=True)
    project = Path(config['project'])
    dotnet = shutil.which('dotnet') or 'C:/Program Files/dotnet/dotnet.exe'
    item = {'release_id': release_id, 'status': 'preparing', 'note': note}
    try:
        subprocess.run([dotnet, 'publish', str(project / settings['project_file']), '-c', 'Release',
                        '--self-contained', 'false', '-p:UseAppHost=false', '-o', str(folder / 'payload')], check=True)
        payload = folder / 'payload'
        files = {p.relative_to(payload).as_posix(): sha(p) for p in payload.rglob('*') if p.is_file()}
        if 'kenton_website.dll' not in files:
            raise ValueError('Published application DLL missing.')
        spec = dict(item, status='prepared', host=settings['host'], ssh_user=settings['ssh_user'],
                    service_name=settings['service_name'], remote_directory=settings['remote_directory'],
                    health_url=settings['health_url'], files=files, preserve_files=settings['preserve_files'])
        revision = subprocess.run(['git', '-C', str(project), 'rev-parse', 'HEAD'], capture_output=True, text=True)
        changes = subprocess.run(['git', '-C', str(project), 'status', '--porcelain'], capture_output=True, text=True)
        spec['source_revision'] = revision.stdout.strip() if revision.returncode == 0 else 'unavailable'
        spec['source_dirty'] = bool(changes.stdout.strip()) if changes.returncode == 0 else None
        for name in spec['preserve_files']:
            if Path(name).name != name:
                raise ValueError('Preserved configuration must be top-level filenames.')
        with tarfile.open(folder / 'payload.tar.gz', 'w:gz') as archive:
            for name in files:
                archive.add(payload / name, arcname=name, recursive=False)
        shutil.copy2(root / 'scripts/_website_remote.py', folder / 'activate.py')
        spec['archive_sha256'] = sha(folder / 'payload.tar.gz')
        spec['activator_sha256'] = sha(folder / 'activate.py')
        (folder / 'release.json').write_text(json.dumps(spec, indent=2), encoding='utf-8')
        history(root, settings, spec)
        print(f'Prepared release: {folder}')
        return folder
    except Exception:
        history(root, settings, dict(item, status='prepare_failed'))
        raise


def log_workbook(root, settings, folder, receipt):
    logging = settings['maintenance_workbook']
    if not logging.get('enabled'):
        return 'disabled'
    import openpyxl
    path = Path(logging['path'])
    original = path.read_bytes()
    backup = folder / 'recent-logs-before.xlsx'
    if not backup.exists():
        backup.write_bytes(original)
    book = openpyxl.load_workbook(path)
    try:
        sheet = book[logging['sheet']] if logging['sheet'] in book.sheetnames else book.create_sheet(logging['sheet'])
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            for column, label in enumerate(['UTC time', 'Release ID', 'URL', 'Status', 'Description', 'Receipt'], 1):
                sheet.cell(1, column, label)
            sheet.freeze_panes = 'A2'
        if any(row[0] == receipt['release_id'] for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)):
            return 'already_logged'
        sheet.append([datetime.now(timezone.utc).isoformat(), receipt['release_id'], settings['public_url'],
                      receipt['status'], receipt.get('note', ''), str(folder / 'receipt.json')])
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as temp:
            staged = Path(temp.name)
        try:
            book.save(staged)
            if path.read_bytes() != original:
                raise ValueError('Workbook changed during logging; retry logging only.')
            staged.replace(path)
        finally:
            staged.unlink(missing_ok=True)
    finally:
        book.close()
    return 'logged'


def deploy(root, config, folder, reviewed):
    settings = config['deployment']
    if not reviewed:
        raise ValueError('Review the prepared release first; deployment requires --reviewed.')
    if not settings.get('enabled') or not settings.get('service_verified'):
        raise ValueError('Deployment is disabled until service and restart permissions are confirmed.')
    folder = Path(folder).resolve()
    if (root / settings['history_directory']).resolve() not in folder.parents:
        raise ValueError('Select a prepared release under the deployment history folder.')
    spec = json.loads((folder / 'release.json').read_text())
    for key in ('host', 'ssh_user', 'service_name', 'remote_directory', 'health_url'):
        if spec[key] != settings[key] or not spec[key]:
            raise ValueError('Release target changed or incomplete: ' + key)
    if urlparse(spec['health_url']).scheme != 'https':
        raise ValueError('A verified HTTPS health URL is required.')
    if sha(folder / 'payload.tar.gz') != spec['archive_sha256'] or sha(folder / 'activate.py') != spec['activator_sha256']:
        raise ValueError('Prepared release changed. Prepare again.')
    if (folder / 'receipt.json').exists():
        raise ValueError('This release has already been attempted; inspect its receipt. Use --log-only for a pending workbook entry.')
    remote = settings['remote_release_root'].rstrip('/') + '/' + spec['release_id']
    if not remote.startswith('/home/') or '..' in Path(remote).parts:
        raise ValueError('Remote release storage must be an explicit path under /home/.')
    host = f"{settings['ssh_user']}@{settings['host']}"
    options = ['-o', 'BatchMode=yes', '-o', 'StrictHostKeyChecking=yes', '-o', 'ConnectTimeout=10']
    if settings.get('identity_file'):
        options += ['-i', str(Path(settings['identity_file']).expanduser())]
    ssh = ['ssh', *options, '-p', str(settings.get('ssh_port', 22)), host]
    receipt = {'release_id': spec['release_id'], 'status': 'deploying', 'note': spec['note'], 'remote_release': remote}
    (folder / 'receipt.json').write_text(json.dumps(receipt, indent=2))
    try:
        subprocess.run(ssh + ['mkdir -p ' + shlex.quote(remote)], check=True)
        subprocess.run(['scp', *options, '-P', str(settings.get('ssh_port', 22)),
                        *[str(folder / n) for n in ('payload.tar.gz', 'release.json', 'activate.py')],
                        host + ':' + shlex.quote(remote) + '/'], check=True)
        result = subprocess.run(ssh + ['python3 ' + shlex.quote(remote + '/activate.py') + ' ' + shlex.quote(remote)],
                                capture_output=True, text=True, timeout=240)
        # Fetch the receipt even when activation rolled back or the connection failed.
        response = subprocess.run(ssh + ['cat ' + shlex.quote(remote + '/remote-receipt.json')],
                                  capture_output=True, text=True, timeout=20)
        if response.returncode == 0:
            receipt.update(json.loads(response.stdout))
        else:
            receipt.update(status='unknown_check_server', error='No remote receipt; inspect before retrying.')
        if result.returncode != 0 and receipt['status'] == 'deployed':
            receipt['transport_warning'] = 'SSH returned an error; remote receipt reports success.'
    except Exception as error:
        receipt.update(status='unknown_check_server', error=str(error))
    (folder / 'receipt.json').write_text(json.dumps(receipt, indent=2))
    history(root, settings, receipt)
    if receipt['status'] != 'deployed':
        raise ValueError(f"Deployment status: {receipt['status']}. Inspect {folder / 'receipt.json'}")
    try:
        receipt['workbook_log'] = log_workbook(root, settings, folder, receipt)
    except Exception as error:
        receipt.update(workbook_log='pending', workbook_error=str(error))
    (folder / 'receipt.json').write_text(json.dumps(receipt, indent=2))
    print(f"Deployment verified. Workbook log: {receipt['workbook_log']}. Receipt: {folder / 'receipt.json'}")
